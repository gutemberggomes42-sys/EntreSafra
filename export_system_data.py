import json
import re
from datetime import date, datetime, time
from pathlib import Path

import openpyxl


SOURCE = Path(r"C:\Users\gutem\Downloads\Controle de manutenções EntreSafra NOVO - Editável.xlsm")
OUTPUT = Path(__file__).with_name("data.json")
JS_OUTPUT = Path(__file__).with_name("data.js")


def clean(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if not isinstance(value, (str, int, float, bool)):
        return str(value)
    if isinstance(value, str):
        return value.replace("\xa0", " ").strip()
    return value


def cell_value(formula_ws, cached_ws, coordinate):
    formula_value = formula_ws[coordinate].value
    cached_value = cached_ws[coordinate].value
    if cached_value not in (None, ""):
        return clean(cached_value)
    return clean(formula_value)


def rows_from_sheet(formula_ws, cached_ws, start, end, columns, headers):
    records = []
    for row in range(start, end + 1):
        record = {
            header: cell_value(formula_ws, cached_ws, f"{column}{row}")
            for column, header in zip(columns, headers)
        }
        if any(value not in ("", " ") for value in record.values()):
            record["_sourceRow"] = row
            records.append(record)
    return records


def raw_sheet(formula_ws, cached_ws):
    rows = []
    for row in formula_ws.iter_rows():
        values = []
        for cell in row:
            value = cell_value(formula_ws, cached_ws, cell.coordinate)
            values.append(value)
        while values and values[-1] == "":
            values.pop()
        if values:
            rows.append({"row": row[0].row, "values": values})
    return {
        "name": formula_ws.title,
        "maxRow": formula_ws.max_row,
        "maxColumn": formula_ws.max_column,
        "rows": rows,
    }


def normalise_status(value):
    text = str(value or "").strip().upper()
    if not text:
        return "Não informado"
    if "FEITO" in text or text in {"OK", "REALIZADO"}:
        return "Concluído"
    if "ANDAMENTO" in text or "MANUTENÇÃO" in text:
        return "Em andamento"
    if "PEND" in text or "AG." in text:
        return "Pendente"
    return str(value).strip()


def main():
    formulas = openpyxl.load_workbook(SOURCE, data_only=False, keep_vba=True)
    cached = openpyxl.load_workbook(SOURCE, data_only=True, keep_vba=True)

    specs = [
        ("radios", "CONTROLE RÁDIOS SETOR AGRÍCOLA", 3, 193, list("ABCDEFG"),
         ["identificador", "frota", "descricao", "setor", "possuiRadio", "carregador", "tipo"]),
        ("plantio", "OFICIAL PLANTIO", 4, 100, list("BCDEFGHIJ"),
         ["equipamento", "modelo", "frota", "frente", "status", "rc", "cdc", "pendencias", "previsaoEntrega"]),
        ("caminhoesReforma", "REFORMA CAMINHOES", 4, 50, list("ABCDEFJKLM"),
         ["frota", "placa", "funcao", "limpeza", "localizacao", "pendencias", "status", "inicio", "fim", "dias"]),
        ("caminhoesInfo", "INFORMAÇÕES CAMINHÕES", 4, 1500, list("ABCDEFGHIJKLMNOPQRS"),
         ["frota", "ano", "placa", "operacao", "placaDianteira", "placaTraseira", "antt", "tacografo", "orcamento", "dataAfericao", "vencimento", "dnit", "der", "possuiCrlv", "anoCrlv", "radio", "adesivo", "tara", "pintura"]),
        ("carretasReforma", "REFORMA CARRETAS", 4, 91, list("ABCDEFGHIJKLMNOPQRS"),
         ["frota", "placa", "grupo", "funcao", "limpezaSeco", "limpezaPipa", "localizacao", "pendencias", "status", "porcasFaltantes", "parafusosFaltantes", "lubrificacao", "obsLubrificacao", "calibracao", "pneus", "inicio", "fim", "dias", "conjunto"]),
        ("carretasInfo", "INFORMAÇÕES CARRETAS", 7, 111, list("ABCDEFGHIJKLMNOP"),
         ["frota", "grupo", "placa", "situacaoPlaca", "lacre", "numeroDianteiro", "numeros", "possuiCrlv", "anoCrlv", "adesivoVeiculoLongo", "antt", "faixaParachoque", "faixaRefletiva", "cordaLonas", "caboSeguranca", "observacao"]),
        ("colhedoras", "REFORMA COLHEDORAS", 4, 14, list("ABCIJKL"),
         ["frota", "localizacao", "pendencias", "status", "inicio", "fim", "dias"]),
        ("tratores", "REFORMA TRATORES", 4, 57, list("ABCDEFG"),
         ["frota", "localizacao", "pendencias", "status", "inicio", "fim", "dias"]),
        ("transbordos", "REFORMA TRANSBORDOS", 4, 38, list("ABCIJKL"),
         ["frota", "localizacao", "pendencias", "status", "inicio", "fim", "dias"]),
        ("vivencias", "REFORMA VIVENCIAS", 4, 9, list("ABCI"),
         ["frota", "localizacao", "pendencias", "status"]),
        ("cci", "REFORMA CCI", 4, 13, list("ABCD"),
         ["frota", "localizacao", "pendencias", "status"]),
        ("curvaS", 'CURVA "S"', 2, 69, list("ABCDEFGHIJKLMN"),
         ["frota", "grupo", "funcao", "planejado", "realizado", "emManutencao", "previsao", "consideracoes", "inicioPlanejado", "fimPlanejado", "diasPlanejados", "entrada", "saida", "diasTrabalhados"]),
    ]

    modules = {}
    for key, sheet_name, start, end, columns, headers in specs:
        modules[key] = rows_from_sheet(
            formulas[sheet_name], cached[sheet_name], start, end, columns, headers
        )
        for index, record in enumerate(modules[key], start=1):
            record["id"] = f"{key}-{index}"
            if "status" in record:
                record["statusNormalizado"] = normalise_status(record["status"])

    axiagro_controle = [
        ["CAMPO","60007","4001","OK","6723","OK","OK","ec:b5:50:73:7a:2d",""],
        ["CAMPO","60008","4001","OK","58298","OK","OK","ec:b5:50:73:81:11",""],
        ["CAMPO","6721","4001","OK","6360","OK","OK","ec:b5:50:73:7a:29",""],
        ["CAMPO","6988","4002","OK","6377","OK","OK","ec:b5:50:73:7a:51",""],
        ["CAMPO","6989","4002","OK","6378","OK","OK","ec:b5:50:73:7a:21",""],
        ["CAMPO","6722","4002","OK","6384","OK","OK","ec:b5:50:73:7a:db",""],
        ["CONTROLE","9828","4003","OK","57843","OK","OK","ec:b5:50:73:7a:1d",""],
        ["CAMPO","9829","4003","OK","58413","OK","OK","ec:b5:50:73:7a:1b",""],
        ["CAMPO","6990","4003","OK","6329","OK","OK","ec:b5:50:73:7a:15",""],
        ["CONTROLE","1082","4002","OK","23889","OK","OK","ec:b5:50:73:7a:2f",""],
        ["CAMPO","60013","4001","era o da 9828","66793","OK","OK","78:37:16:ae:4c:64","Substituição"],
        ["CAMPO","60009","4001","era o da 60008","66791","OK","OK","78:37:16:ae:4b:5c","Substituição"],
        ["CAMPO","60010","4001","era o da 6721","67790","OK","OK","78:37:16:ae:57:be","Substituição"],
        ["CONTROLE","6988","4002","OK","","OK","OK","78:37:16:ae:54:96",""],
        ["CAMPO","6611 (novo)","4001","verificar mac do antigo 6611","","OK","OK","78:37:16:ae:4d:d2","Verificar MAC"],
        ["CAMPO","6516 (novo)","4002","era o da 6722","","OK","OK","78:37:16:ae:4d:b8","Substituição"],
        ["CAMPO","9829","4003","OK","","OK","OK","78:37:16:ae:3d:92",""],
        ["CAMPO","6964 (novo)","4099","era o da 6990","","OK","OK","78:37:16:ad:aa:fc","Substituição"],
        ["CAMPO","6668 (novo)","4003","era o da 1082","","OK","OK","78:37:16:ae:55:c8","Substituição"],
        ["CONTROLE","6611","4001","OK","6359","X","OK","78:37:16:ae:3f:a8","Fusível pendente"],
        ["CONTROLE","6612","4001","OK","6372","OK","OK","78:37:16:ae:ab:08",""],
        ["CONTROLE","6617 - foi p/T.I","4001","OK","6396","OK","OK","78:37:16:ae:a5:40","Foi para T.I"],
        ["6526 transformado para (6528 novo)","6528","4001","OK","6309","OK","OK","78:37:16:ae:7e:b8","Transformação de frota"],
        ["CONTROLE","6616 - foi p/T.I","4001","OK","6359","OK","OK","78:37:16:ae:59:5c","Foi para T.I"],
        ["6528 transformada para (6965 novo)","6965 caminhão","4099","OK","6391","OK","OK","78:37:16:ae:3d:aa","Transformação de frota"],
        ["CONTROLE","6527","4002","OK","23879","OK","OK","78:37:16:ae:4f:24",""],
        ["CONTROLE","6969","4001","OK","57665","OK","OK","78:37:16:ae:7c:b8",""],
        ["CONTROLE","6966","4002","OK","6333","OK","OK","78:37:16:ae:5a:36",""],
        ["CONTROLE","6967","4002","OK","6304","X","OK","78:37:16:ae:aa:fe","Fusível pendente"],
        ["CONTROLE","6968","4002","OK","6316","OK","OK","78:37:16:ae:58:c6",""],
        ["Telefone não encontrado - feito b.o","6668","4003","OK","6327","OK","OK","78:37:16:ae:4a:7a","Telefone não encontrado"],
        ["CONTROLE","6669","4003","OK","6385","OK","OK","78:37:16:ae:3d:ca",""],
        ["CONTROLE","6670","4003","OK","6331","","OK","78:37:16:ae:4c:a6",""],
        ["CONTROLE","6615","4003","Novo 6995 N1","7948","OK","OK","78:37:16:ae:a5:58","Substituição"],
        ["NO CONTROLE","6525","***","REFORMA","-","","OK","78:37:16:ae:7e:c0","Em reforma"],
        ["NO CONTROLE","6458","***","OK","57662","","-","","Sem controle"],
        ["NO CONTROLE","6965 Novo","4099","OK","6355","","OK","78:37:16:ae:a5:40","Substituído/6963"],
        ["ATIVO","6963 novo","4099","OK","58305","","OK","78:37:16:ae:3d:b8","Substituído/6682"],
        ["ATIVO","6698","4099","OK","6319","","OK","78:37:16:ae:4d:82",""],
        ["ATIVO","6699","4099","OK","6352","","OK","78:37:16:ae:54:7c",""],
        ["ATIVO","6700","4099","OK","6400","","OK","78:37:16:ae:31:28",""],
        ["ATIVO","6960","4099","OK","6703","","OK","78:37:16:ae:54:3a",""],
        ["ATIVO","6961","4099","OK","6336","","OK","78:37:16:ae:5b:44",""],
        ["ATIVO","6962","4099","OK","6305","","OK","78:37:16:ae:a5:70","Substituído para 4786"],
        ["ATIVO","6965","4099","OK","6332","","OK","78:37:16:ae:4c:64","Substituído para 9828"],
        ["NO CONTROLE","6964","4099","OK","67786","","OK","78:37:16:ae:30:2e","Substituído para 6416"],
        ["ATIVO","6964 (novo)","4100","OK","66595","","OK","78:37:16:ad:aa:fc","Era o da colhedora 6990"],
        ["NO CONTROLE","6965","4099","OK","6718","","OK","78:37:16:ad:aa:f6","Substituído/6680"],
        ["ATIVO","6995","4099","OK","67787","","OK","78:37:16:ae:7e:74",""],
        ["ATIVO","6996","4099","OK","6322","","OK","78:37:16:ae:7f:02",""],
        ["ATIVO","6997","4099","OK","6363","","OK","78:37:16:ad:ab:02",""],
        ["ATIVO","6998","4099","OK","67288","","OK","78:37:16:ae:4c:92","Substituído para 6150; agora 6650"],
        ["ATIVO","6999","4099","OK","6342","","OK","78:37:16:ae:7e:c0",""],
        ["","6650","4099","OK","","","","d8:cf:bf:2f:85:f7",""],
        ["ATIVO","6682","4099","OK","6362","","OK","78:37:16:ae:3f:a8",""],
        ["ATIVO","6995","4099","OK","","","OK","78:37:16:ae:a5:58","Substituído/6615"],
        ["ATIVO","6965","4099","OK","66528","","OK","78:37:16:ae:3d:aa",""],
    ]
    modules["axiagroControle"] = [
        {"statusCelular": row[0], "frota": row[1], "frente": row[2], "statusSuporte": row[3], "numeroLacre": row[4], "fusivel": row[5], "statusAparelho": row[6], "mac": row[7], "observacao": row[8], "id": f"axiagroControle-{index}"}
        for index, row in enumerate(axiagro_controle, start=1)
    ]
    axiagro_estoque = [
        ["Suportes Carregadores","C-20","-"], ["Carregador de Carro USB","LE-6821",12],
        ["Cabos carregador V8","V8 -35W",16], ["Cabos carregador - Tipo C","MD-8440 tipo - C",26],
        ["Fita Isolante","P22",0], ["Fusível de lâmina","10A",1],
        ["Fita Crepe verde (NORTON)","18mm x 40m","-"], ["Fita Transparente (FIT - PEL)","",2],
        ["Fita Adesiva","45mm x 45m",4], ["Antena starlink","starlink mini",10],
        ["Quick Charger USB","55W",2], ["Tomada de força DNI automotivo","12/24 V",18],
        ["Auto Cigarette","acendedor de cigarro",8], ["Suporte para PTT","-","-"],
        ["Antena kit starlink","-",8], ["Controladoras","XTRA4415N",2],
        ["Controladoras","XTRA4415N-XDS2",2],
    ]
    modules["axiagroEstoque"] = [
        {"equipamento": row[0], "modelo": row[1], "quantidade": row[2], "id": f"axiagroEstoque-{index}"}
        for index, row in enumerate(axiagro_estoque, start=1)
    ]
    employee_source = Path(__file__).with_name("funcionarios_import.json")
    if employee_source.exists():
        employee_data = json.loads(employee_source.read_text(encoding="utf-8"))
        modules["funcionarios"] = employee_data["records"]

    raw = [raw_sheet(ws, cached[ws.title]) for ws in formulas.worksheets]
    result = {
        "meta": {
            "title": "Gestão de Manutenções EntreSafra",
            "sourceFile": SOURCE.name,
            "importedAt": datetime.now().isoformat(timespec="seconds"),
            "sheetCount": len(formulas.sheetnames),
            "sourceSheets": formulas.sheetnames,
        },
        "modules": modules,
        "rawSheets": raw,
    }
    compact = json.dumps(result, ensure_ascii=False, separators=(",", ":"))
    OUTPUT.write_text(compact, encoding="utf-8")
    JS_OUTPUT.write_text("window.__ENTRESSAFRA_DATA__=" + compact + ";", encoding="utf-8")
    print(f"{OUTPUT} ({OUTPUT.stat().st_size} bytes)")
    print(f"{JS_OUTPUT} ({JS_OUTPUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
